from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
import re
from typing import BinaryIO

from openpyxl import Workbook, load_workbook

from .models import (
    DemandRecord,
    normalize_text,
    parse_credit_value_to_cent,
    validate_credit_type,
)


REQUIRED_HEADERS = {
    "学号": "student_id",
    "姓名": "student_name",
    "学分类型": "credit_type",
    "学分数值": "credit_value",
}


@dataclass(frozen=True, slots=True)
class PriorityDemandRecord:
    student_id: str
    student_name: str
    credit_types: tuple[str, ...]
    requested_value_cent: int


def _resolve_workbook(source: Workbook | str | Path | BinaryIO) -> Workbook:
    if isinstance(source, Workbook):
        return source
    return load_workbook(source)


def _detect_columns(header_row: tuple[object, ...]) -> dict[str, int]:
    columns = {}
    for index, value in enumerate(header_row):
        header = normalize_text(value)
        if header in REQUIRED_HEADERS:
            columns[REQUIRED_HEADERS[header]] = index

    missing_headers = set(REQUIRED_HEADERS.values()) - set(columns)
    if missing_headers:
        raise ValueError(f"Missing required headers: {sorted(missing_headers)}")

    return columns


def _parse_credit_types(value: str) -> tuple[str, ...]:
    credit_types = []
    for item in re.split(r"[,，]", value):
        credit_type_text = normalize_text(item)
        if not credit_type_text:
            continue
        credit_type = validate_credit_type(credit_type_text)
        if credit_type not in credit_types:
            credit_types.append(credit_type)

    if not credit_types:
        raise ValueError("Credit type is required")

    return tuple(credit_types)


def parse_priority_input_workbook(
    source: Workbook | str | Path | BinaryIO,
) -> list[PriorityDemandRecord]:
    workbook = _resolve_workbook(source)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    columns = _detect_columns(rows[0])
    parsed_rows: list[PriorityDemandRecord] = []

    for row in rows[1:]:
        student_id = normalize_text(row[columns["student_id"]])
        student_name = normalize_text(row[columns["student_name"]])
        credit_type_text = normalize_text(row[columns["credit_type"]])
        credit_value = row[columns["credit_value"]]

        if not any((student_id, student_name, credit_type_text, normalize_text(credit_value))):
            continue

        credit_types = _parse_credit_types(credit_type_text)
        requested_value_cent = parse_credit_value_to_cent(credit_value)

        if not student_id or not student_name:
            raise ValueError("Student id and name are required")

        parsed_rows.append(
            PriorityDemandRecord(
                student_id=student_id,
                student_name=student_name,
                credit_types=credit_types,
                requested_value_cent=requested_value_cent,
            )
        )

    return parsed_rows


def parse_input_workbook(source: Workbook | str | Path | BinaryIO) -> list[DemandRecord]:
    parsed_rows: list[DemandRecord] = []
    for row in parse_priority_input_workbook(source):
        if len(row.credit_types) != 1:
            raise ValueError(
                "Multiple credit types require priority demand planning"
            )
        parsed_rows.append(
            DemandRecord(
                student_id=row.student_id,
                student_name=row.student_name,
                credit_type=row.credit_types[0],
                requested_value_cent=row.requested_value_cent,
            )
        )
    return parsed_rows


def load_and_aggregate_demands(
    source: Workbook | str | Path | BinaryIO,
) -> list[DemandRecord]:
    aggregated_values: dict[tuple[str, str], int] = defaultdict(int)
    student_names: dict[tuple[str, str], str] = {}

    for row in parse_input_workbook(source):
        key = (row.student_id, row.credit_type)
        aggregated_values[key] += row.requested_value_cent
        student_names.setdefault(key, row.student_name)

    return [
        DemandRecord(
            student_id=student_id,
            student_name=student_names[(student_id, credit_type)],
            credit_type=credit_type,
            requested_value_cent=aggregated_values[(student_id, credit_type)],
        )
        for student_id, credit_type in sorted(aggregated_values)
    ]


def load_and_aggregate_priority_demands(
    source: Workbook | str | Path | BinaryIO,
) -> list[PriorityDemandRecord]:
    aggregated_values: dict[tuple[str, tuple[str, ...]], int] = {}
    student_names: dict[tuple[str, tuple[str, ...]], str] = {}

    for row in parse_priority_input_workbook(source):
        key = (row.student_id, row.credit_types)
        aggregated_values[key] = aggregated_values.get(key, 0) + row.requested_value_cent
        student_names.setdefault(key, row.student_name)

    return [
        PriorityDemandRecord(
            student_id=student_id,
            student_name=student_names[(student_id, credit_types)],
            credit_types=credit_types,
            requested_value_cent=requested_value_cent,
        )
        for (student_id, credit_types), requested_value_cent in aggregated_values.items()
    ]
