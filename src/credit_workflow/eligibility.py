from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .models import ActivityBundle, DemandRecord, normalize_text


@dataclass(frozen=True, slots=True)
class AdmitMember:
    sign_up_id: str
    student_name: str
    student_id: str
    user_id: str


@dataclass(frozen=True, slots=True)
class EligibleBundleMatch:
    activity_id: str
    activity_name: str
    credit_type: str
    sign_up_id: str
    user_id: str


@dataclass(slots=True)
class EligibilityResult:
    matches_by_demand: dict[tuple[str, str, str], tuple[EligibleBundleMatch, ...]]
    not_in_admit_list: tuple[DemandRecord, ...]

    def matches_for_demand(self, demand: DemandRecord) -> tuple[EligibleBundleMatch, ...]:
        return self.matches_by_demand.get(
            (demand.student_id, demand.student_name, demand.credit_type), ()
        )

    def is_eligible(self, student_id: str, activity_id: str, credit_type: str) -> bool:
        return any(
            match.activity_id == activity_id and match.credit_type == credit_type
            for (candidate_student_id, _, candidate_credit_type), matches in self.matches_by_demand.items()
            if candidate_student_id == student_id and candidate_credit_type == credit_type
            for match in matches
        )

    def get_match(
        self,
        student_id: str,
        student_name: str,
        credit_type: str,
        activity_id: str,
    ) -> EligibleBundleMatch | None:
        for match in self.matches_by_demand.get(
            (student_id, student_name, credit_type), ()
        ):
            if match.activity_id == activity_id:
                return match
        return None


def _get_text(row: tuple[object, ...], index: int, default: str = "") -> str:
    if index >= len(row):
        return default
    return normalize_text(row[index])


def _load_admit_workbook(content: bytes):
    return load_workbook(BytesIO(content))


def _normalize_header(value: object) -> str:
    return "".join(ch for ch in normalize_text(value).lower() if ch.isalnum())


def _resolve_admit_column_indexes(header_row: tuple[object, ...]) -> tuple[int, int, int, int]:
    normalized_headers = [_normalize_header(value) for value in header_row]

    def _find_exact(candidates: tuple[str, ...]) -> int | None:
        for candidate in candidates:
            candidate_key = _normalize_header(candidate)
            for index, header in enumerate(normalized_headers):
                if header == candidate_key:
                    return index
        return None

    sign_up_id_index = _find_exact(("报名ID", "报名编号", "signUpId", "signupId"))
    student_name_index = _find_exact(("姓名", "学生姓名", "name", "studentName"))
    student_id_index = _find_exact(("学号", "学生学号", "studentId", "studentNo"))

    user_id_index = _find_exact(
        ("userId", "uid", "用户ID", "用户UID", "用户编号", "userId(uid)")
    )
    if user_id_index is None:
        for index, header in enumerate(normalized_headers):
            if (
                ("userid" in header or "uid" in header or "用户id" in header or "用户uid" in header)
                and "signup" not in header
                and "报名" not in header
            ):
                user_id_index = index
                break

    return (
        sign_up_id_index if sign_up_id_index is not None else 0,
        student_name_index if student_name_index is not None else 3,
        student_id_index if student_id_index is not None else 4,
        user_id_index if user_id_index is not None else 5,
    )


def _parse_admit_members_from_workbook(workbook) -> tuple[AdmitMember, ...]:
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return ()

    sign_up_id_index, student_name_index, student_id_index, user_id_index = (
        _resolve_admit_column_indexes(rows[0])
    )

    members: list[AdmitMember] = []
    for row in rows[1:]:
        sign_up_id = _get_text(row, sign_up_id_index)
        student_name = _get_text(row, student_name_index)
        student_id = _get_text(row, student_id_index)
        user_id = _get_text(row, user_id_index)

        if not any((sign_up_id, student_name, student_id, user_id)):
            continue

        if not sign_up_id or not student_name or not student_id:
            raise ValueError("Admit list row is missing required fields")

        members.append(
            AdmitMember(
                sign_up_id=sign_up_id,
                student_name=student_name,
                student_id=student_id,
                user_id=user_id,
            )
        )

    return tuple(members)


def parse_admit_members(content: bytes) -> tuple[AdmitMember, ...]:
    return _parse_admit_members_from_workbook(_load_admit_workbook(content))


def _load_admit_members_with_retries(
    dmapi, activity_id: str, retry_attempts: int
) -> tuple[AdmitMember, ...]:
    last_error: Exception | None = None
    for _ in range(retry_attempts):
        try:
            content = dmapi.export_mem_excel(activity_id, dmapi.EXPORT_TYPE_ADMIT)
        except Exception as exc:  # pragma: no cover - exercised by retry tests
            last_error = exc
            continue

        try:
            if not content:
                raise ValueError("Admit list export returned no content")
            workbook = _load_admit_workbook(content)
        except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
            last_error = exc
            continue

        return _parse_admit_members_from_workbook(workbook)

    if last_error is not None:
        raise last_error
    return ()


def build_eligibility_map(
    dmapi,
    bundles: list[ActivityBundle],
    demands: list[DemandRecord],
    retry_attempts: int = 3,
) -> EligibilityResult:
    members_by_activity: dict[str, dict[tuple[str, str], AdmitMember]] = {}
    # Keep admit exports deterministic so activity-scoped results never depend on
    # hash/set iteration order.
    for activity_id in sorted({bundle.activity_id for bundle in bundles}):
        members = _load_admit_members_with_retries(dmapi, activity_id, retry_attempts)
        members_by_activity[activity_id] = {
            (member.student_id, member.student_name): member for member in members
        }

    grouped_matches: dict[tuple[str, str, str], list[EligibleBundleMatch]] = defaultdict(list)
    not_in_admit_list: list[DemandRecord] = []

    for demand in demands:
        demand_key = (demand.student_id, demand.student_name, demand.credit_type)
        for bundle in bundles:
            if bundle.credit_type != demand.credit_type:
                continue

            member = members_by_activity[bundle.activity_id].get(
                (demand.student_id, demand.student_name)
            )
            if member is None:
                continue

            grouped_matches[demand_key].append(
                EligibleBundleMatch(
                    activity_id=bundle.activity_id,
                    activity_name=bundle.activity_name,
                    credit_type=bundle.credit_type,
                    sign_up_id=member.sign_up_id,
                    user_id=member.user_id,
                )
            )

        if not grouped_matches[demand_key]:
            not_in_admit_list.append(demand)

    return EligibilityResult(
        matches_by_demand={
            key: tuple(dict.fromkeys(matches))
            for key, matches in grouped_matches.items()
        },
        not_in_admit_list=tuple(not_in_admit_list),
    )
