from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from .eligibility import EligibilityResult
from .models import ActivityBundle, DemandRecord
from .planner import DemandAllocation, PlanningResult, bundle_key


def activity_directory_name(activity_id: str, activity_name: str) -> str:
    safe_name = activity_name.replace("/", "_").replace("\\", "_")
    return f"{activity_id}_{safe_name}"


def _write_workbook(path: Path, headers: tuple[str, ...], rows: list[tuple[object, ...]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def _demand_rows(demands: tuple[DemandRecord, ...]) -> list[tuple[object, ...]]:
    return [
        (
            demand.student_id,
            demand.student_name,
            demand.credit_type,
            demand.requested_value_cent,
        )
        for demand in demands
    ]


def _allocation_rows(allocations: tuple[DemandAllocation, ...]) -> list[tuple[object, ...]]:
    rows = []
    for allocation in allocations:
        rows.append(
            (
                allocation.demand.student_id,
                allocation.demand.student_name,
                allocation.demand.credit_type,
                allocation.demand.requested_value_cent,
                allocation.planned_value_cent,
                allocation.delta_cent,
                ",".join(
                    assignment.bundle.activity_id for assignment in allocation.assignments
                ),
            )
        )
    return rows


def _bundle_rows(bundles: tuple[ActivityBundle, ...]) -> list[tuple[object, ...]]:
    rows = []
    for bundle in bundles:
        rows.append(
            (
                bundle.activity_id,
                bundle.activity_name,
                bundle.credit_type,
                bundle.bundle_value_cent,
                bundle.bundle_capacity,
                ",".join(item.credit_id for item in bundle.credit_items),
            )
        )
    return rows


def _bundle_usage_rows(
    bundles: tuple[ActivityBundle, ...], planning_result: PlanningResult
) -> list[tuple[object, ...]]:
    rows = []
    for bundle in bundles:
        key = bundle_key(bundle)
        rows.append(
            (
                bundle.activity_id,
                bundle.activity_name,
                bundle.credit_type,
                planning_result.bundle_assignment_counts.get(key, 0),
                bundle.bundle_capacity,
            )
        )
    return rows


def _activity_credit_types(
    bundles: tuple[ActivityBundle, ...],
) -> dict[tuple[str, str], set[str]]:
    activity_credit_types: dict[tuple[str, str], set[str]] = {}
    for bundle in bundles:
        key = (bundle.activity_id, bundle.activity_name)
        activity_credit_types.setdefault(key, set()).add(bundle.credit_type)
    return activity_credit_types


def _write_plan_summary(
    path: Path,
    demands: tuple[DemandRecord, ...],
    bundles: tuple[ActivityBundle, ...],
    planning_result: PlanningResult,
) -> None:
    workbook = Workbook()
    sheets = [
        (
            "需求汇总",
            ("学号", "姓名", "学分类型", "应发值(cent)"),
            _demand_rows(demands),
        ),
        (
            "活动资源",
            ("activityId", "活动名称", "学分类型", "活动包值(cent)", "容量", "creditIds"),
            _bundle_rows(bundles),
        ),
        (
            "学生分配计划",
            ("学号", "姓名", "学分类型", "应发值(cent)", "计划值(cent)", "偏差(cent)", "活动列表"),
            _allocation_rows(planning_result.allocations),
        ),
        (
            "活动均衡统计",
            ("activityId", "活动名称", "学分类型", "已规划人数", "容量"),
            _bundle_usage_rows(bundles, planning_result),
        ),
    ]

    first_sheet = workbook.active
    first_sheet.title = sheets[0][0]
    for sheet_index, (title, headers, rows) in enumerate(sheets):
        sheet = first_sheet if sheet_index == 0 else workbook.create_sheet(title=title)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)

    workbook.save(path)


def write_plan_reports(
    batch_dir: Path,
    demands: tuple[DemandRecord, ...],
    bundles: tuple[ActivityBundle, ...],
    eligibility: EligibilityResult,
    planning_result: PlanningResult,
    preissued_preview: dict[str, tuple[dict, ...]],
) -> None:
    batch_dir.mkdir(parents=True, exist_ok=True)
    activities_dir = batch_dir / "activities"
    activities_dir.mkdir(exist_ok=True)

    not_in_admit_set = {
        (demand.student_id, demand.student_name, demand.credit_type)
        for demand in eligibility.not_in_admit_list
    }

    over_issued_rows = [
        row
        for allocation, row in zip(
            planning_result.allocations, _allocation_rows(planning_result.allocations)
        )
        if allocation.delta_cent > 0
        and (
            allocation.demand.student_id,
            allocation.demand.student_name,
            allocation.demand.credit_type,
        )
        not in not_in_admit_set
    ]
    under_issued_rows = [
        row
        for allocation, row in zip(
            planning_result.allocations, _allocation_rows(planning_result.allocations)
        )
        if allocation.delta_cent < 0
        and (
            allocation.demand.student_id,
            allocation.demand.student_name,
            allocation.demand.credit_type,
        )
        not in not_in_admit_set
    ]

    _write_plan_summary(batch_dir / "plan_summary.xlsx", demands, bundles, planning_result)
    _write_workbook(
        batch_dir / "activities.xlsx",
        ("activityId", "活动名称", "学分类型", "活动包值(cent)", "容量", "creditIds"),
        _bundle_rows(bundles),
    )
    _write_workbook(
        batch_dir / "not_in_admit_list.xlsx",
        ("学号", "姓名", "学分类型", "应发值(cent)"),
        _demand_rows(eligibility.not_in_admit_list),
    )
    _write_workbook(
        batch_dir / "over_issued.xlsx",
        ("学号", "姓名", "学分类型", "应发值(cent)", "计划值(cent)", "偏差(cent)", "活动列表"),
        over_issued_rows,
    )
    _write_workbook(
        batch_dir / "under_issued.xlsx",
        ("学号", "姓名", "学分类型", "应发值(cent)", "计划值(cent)", "偏差(cent)", "活动列表"),
        under_issued_rows,
    )
    _write_workbook(
        batch_dir / "already_partially_issued.xlsx",
        ("学号", "姓名", "学分类型", "说明"),
        _preissued_rows(preissued_preview["already_partially_issued"]),
    )
    _write_workbook(
        batch_dir / "already_fully_issued.xlsx",
        ("学号", "姓名", "学分类型", "说明"),
        _preissued_rows(preissued_preview["already_fully_issued"]),
    )

    for (activity_id, activity_name), credit_types in _activity_credit_types(
        bundles
    ).items():
        activity_dir = activities_dir / activity_directory_name(
            activity_id, activity_name
        )
        activity_dir.mkdir(parents=True, exist_ok=True)

        per_activity_not_in_admit = [
            (
                demand.student_id,
                demand.student_name,
                demand.credit_type,
                demand.requested_value_cent,
            )
            for demand in demands
            if demand.credit_type in credit_types
            and eligibility.get_match(
                demand.student_id,
                demand.student_name,
                demand.credit_type,
                activity_id,
            )
            is None
        ]
        per_activity_over = [
            row
            for allocation, row in zip(
                planning_result.allocations, _allocation_rows(planning_result.allocations)
            )
            if allocation.delta_cent > 0
            and any(
                assignment.bundle.activity_id == activity_id
                for assignment in allocation.assignments
            )
        ]
        per_activity_under = [
            row
            for allocation, row in zip(
                planning_result.allocations, _allocation_rows(planning_result.allocations)
            )
            if allocation.delta_cent < 0
            and any(
                assignment.bundle.activity_id == activity_id
                for assignment in allocation.assignments
            )
        ]
        per_activity_partial = [
            record
            for record in preissued_preview["already_partially_issued"]
            if record["activity_id"] == activity_id
        ]
        per_activity_full = [
            record
            for record in preissued_preview["already_fully_issued"]
            if record["activity_id"] == activity_id
        ]

        _write_workbook(
            activity_dir / "not_in_admit_list.xlsx",
            ("学号", "姓名", "学分类型", "应发值(cent)"),
            per_activity_not_in_admit,
        )
        _write_workbook(
            activity_dir / "over_issued.xlsx",
            ("学号", "姓名", "学分类型", "应发值(cent)", "计划值(cent)", "偏差(cent)", "活动列表"),
            per_activity_over,
        )
        _write_workbook(
            activity_dir / "under_issued.xlsx",
            ("学号", "姓名", "学分类型", "应发值(cent)", "计划值(cent)", "偏差(cent)", "活动列表"),
            per_activity_under,
        )
        _write_workbook(
            activity_dir / "already_partially_issued.xlsx",
            ("学号", "姓名", "学分类型", "说明"),
            _preissued_rows(per_activity_partial),
        )
        _write_workbook(
            activity_dir / "already_fully_issued.xlsx",
            ("学号", "姓名", "学分类型", "说明"),
            _preissued_rows(per_activity_full),
        )
        _write_workbook(
            activity_dir / "final_failures.xlsx",
            ("学号", "姓名", "学分类型", "说明"),
            [],
        )


def _execution_rows(records: tuple[dict, ...] | list[dict]) -> list[tuple[object, ...]]:
    return [
        (
            record["activity_id"],
            record["activity_name"],
            record["credit_type"],
            record["student_id"],
            record["student_name"],
            record["sign_up_id"],
            record["user_id"],
            record["credit_id"],
            record["note"],
        )
        for record in records
    ]


def _preissued_rows(records: tuple[dict, ...] | list[dict]) -> list[tuple[object, ...]]:
    return [
        (
            record["student_id"],
            record["student_name"],
            record["credit_type"],
            record["note"],
        )
        for record in records
    ]


def _write_execution_summary(
    path: Path,
    execution_payload: dict[str, tuple[dict, ...]],
) -> None:
    workbook = Workbook()
    sheets = [
        ("补签成功", execution_payload["resign_successes"]),
        ("发放成功", execution_payload["issue_successes"]),
        ("执行期不在录取名单", execution_payload["admit_conflicts"]),
        ("执行期容量冲突", execution_payload["capacity_conflicts"]),
        ("已部分发放", execution_payload["already_partially_issued"]),
        ("已全部发放", execution_payload["already_fully_issued"]),
        ("发放最终失败", execution_payload["final_failures"]),
    ]
    headers = (
        "activityId",
        "活动名称",
        "学分类型",
        "学号",
        "姓名",
        "signUpId",
        "userId",
        "creditId",
        "说明",
    )

    first_sheet = workbook.active
    first_sheet.title = sheets[0][0]
    for sheet_index, (title, records) in enumerate(sheets):
        sheet = first_sheet if sheet_index == 0 else workbook.create_sheet(title=title)
        sheet.append(headers)
        for row in _execution_rows(records):
            sheet.append(row)

    workbook.save(path)


def write_execution_reports(
    batch_dir: Path,
    execution_payload: dict[str, tuple[dict, ...]],
    final_student_summary_rows: list[tuple[object, ...]],
) -> None:
    batch_dir.mkdir(parents=True, exist_ok=True)
    activities_dir = batch_dir / "activities"
    activities_dir.mkdir(exist_ok=True)

    _write_execution_summary(batch_dir / "execution_results.xlsx", execution_payload)
    _write_workbook(
        batch_dir / "retry_results.xlsx",
        (
            "activityId",
            "活动名称",
            "学分类型",
            "学号",
            "姓名",
            "signUpId",
            "userId",
            "creditId",
            "说明",
        ),
        _execution_rows(execution_payload["retry_successes"]),
    )
    _write_workbook(
        batch_dir / "final_student_summary.xlsx",
        ("学号", "姓名", "学分类型", "应发值(cent)", "实际发放值(cent)", "最终偏差(cent)", "状态"),
        final_student_summary_rows,
    )

    activity_records: dict[tuple[str, str], dict[str, list[dict]]] = {}
    per_activity_keys = (
        "admit_conflicts",
        "capacity_conflicts",
        "already_partially_issued",
        "already_fully_issued",
        "final_failures",
    )
    for record_type in per_activity_keys:
        for record in execution_payload[record_type]:
            key = (record["activity_id"], record["activity_name"])
            activity_records.setdefault(
                key,
                {name: [] for name in per_activity_keys},
            )[record_type].append(record)

    for (activity_id, activity_name), records_by_type in activity_records.items():
        activity_dir = activities_dir / activity_directory_name(activity_id, activity_name)
        activity_dir.mkdir(parents=True, exist_ok=True)
        for record_type in per_activity_keys:
            _write_workbook(
                activity_dir / f"{record_type}.xlsx",
                (
                    "activityId",
                    "活动名称",
                    "学分类型",
                    "学号",
                    "姓名",
                    "signUpId",
                    "userId",
                    "creditId",
                    "说明",
                ),
                _execution_rows(records_by_type[record_type]),
            )
