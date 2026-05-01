from openpyxl import load_workbook

from random_drain.models import CandidateMember, CreditDrainSelection, DrainRecord
from random_drain.reports import write_random_drain_reports


def test_write_random_drain_reports(tmp_path):
    member = CandidateMember("9001", "u-001", "张三")
    selection = CreditDrainSelection(
        activity_id="1001",
        activity_name="活动A",
        credit_id="201",
        credit_type="美育实践学分",
        total_capacity=100,
        provided_count=70,
        threshold=90,
        selected_members=(member,),
        candidate_count=1,
        status="ready",
        threshold_percent=90,
        base_target_count=90,
        jitter_offset=0,
        final_target_count=90,
    )
    success = DrainRecord(
        "1001",
        "活动A",
        "美育实践学分",
        "201",
        "9001",
        "u-001",
        "张三",
        "Credit issued successfully",
    )

    write_random_drain_reports(
        tmp_path,
        selections=(selection,),
        successes=(success,),
        failures=(),
        skipped=(),
    )

    assert (tmp_path / "random_drain_summary.xlsx").exists()
    assert (tmp_path / "random_drain_successes.xlsx").exists()
    rows = list(
        load_workbook(tmp_path / "random_drain_summary.xlsx")
        .active.iter_rows(values_only=True)
    )
    assert rows[0] == (
        "activityId",
        "活动名称",
        "学分类型",
        "creditId",
        "总名额",
        "执行前已发放",
        "阈值百分比",
        "基础目标人数",
        "随机偏移人数",
        "最终目标人数",
        "候选人数",
        "本次计划发放",
        "状态",
        "说明",
    )
    assert rows[1][0:5] == ("1001", "活动A", "美育实践学分", "201", 100)
