from __future__ import annotations

import json
import re
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from credit_workflow.models import ActivityBundle, CreditItem, DemandRecord
from credit_workflow.eligibility import build_eligibility_map
from credit_workflow.workflow import run_plan_workflow
import main as cli_main


def build_admit_workbook_bytes(
    rows: list[list[object]], header: list[object] | None = None
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    if header is None:
        header = ["报名ID", "活动名称", "状态", "姓名", "学号", "userId"]
    sheet.append(header)
    for row in rows:
        sheet.append(row)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_bundle(activity_id: str = "1001") -> ActivityBundle:
    return ActivityBundle(
        activity_id=activity_id,
        activity_name="美育活动A",
        credit_type="美育实践学分",
        bundle_value_cent=50,
        bundle_capacity=9,
        credit_items=(
            CreditItem(
                credit_id="201",
                score_id="301",
                credit_type="美育实践学分",
                unitcount_cent=20,
                remaining_capacity=12,
            ),
            CreditItem(
                credit_id="202",
                score_id="301",
                credit_type="美育实践学分",
                unitcount_cent=30,
                remaining_capacity=9,
            ),
        ),
    )


class FakeDMAPI:
    EXPORT_TYPE_ADMIT = 2

    def __init__(self, export_results: list[bytes | None]):
        self.export_results = list(export_results)
        self.export_calls: list[tuple[str, int]] = []

    def export_mem_excel(self, activity_id: str, export_type: int) -> bytes | None:
        self.export_calls.append((activity_id, export_type))
        if not self.export_results:
            raise AssertionError("No more fake export results configured")
        return self.export_results.pop(0)


class FakePlanningDMAPI(FakeDMAPI):
    CREDIT_URL_CREDITEDMEM = "credited"

    def get_mime_manage_activity_dict(self):
        return {
            "1001": {"activityId": 1001, "name": "美育活动A"},
            "1002": {"activityId": 1002, "name": "美育活动B"},
        }

    def get_creditType_list(self, activity_id: str):
        credit_rows = {
            "1001": [
                {
                    "creditId": 201,
                    "scoreId": 301,
                    "scorename": "美育实践学分",
                    "scoretypename": "学分",
                    "unitcount": "0.20",
                    "num": 20,
                    "providenum": 0,
                    "finish": False,
                },
                {
                    "creditId": 202,
                    "scoreId": 301,
                    "scorename": "美育实践学分",
                    "scoretypename": "学分",
                    "unitcount": "0.30",
                    "num": 20,
                    "providenum": 0,
                    "finish": False,
                },
            ],
            "1002": [
                {
                    "creditId": 203,
                    "scoreId": 302,
                    "scorename": "美育实践学分",
                    "scoretypename": "学分",
                    "unitcount": "0.50",
                    "num": 20,
                    "providenum": 0,
                    "finish": False,
                }
            ],
        }
        return credit_rows[str(activity_id)]

    def get_credit_list(self, url: str, activity_id: str, score_id: str):
        assert url == self.CREDIT_URL_CREDITEDMEM
        return []


class SignCardPlanningDMAPI(FakePlanningDMAPI):
    def __init__(self, export_results: list[bytes | None], sign_ids: dict[str, str | None]):
        super().__init__(export_results)
        self.sign_ids = {str(key): value for key, value in sign_ids.items()}
        self.sign_id_calls: list[str] = []
        self.credit_type_calls: list[str] = []

    def get_signid(self, activity_id: str):
        self.sign_id_calls.append(str(activity_id))
        return self.sign_ids.get(str(activity_id))

    def get_creditType_list(self, activity_id: str):
        self.credit_type_calls.append(str(activity_id))
        return super().get_creditType_list(activity_id)


class FakeMultiBundlePlanningDMAPI(FakeDMAPI):
    def get_mime_manage_activity_dict(self):
        return {
            "1001": {"activityId": 1001, "name": "综合活动A"},
        }

    def get_creditType_list(self, activity_id: str):
        assert str(activity_id) == "1001"
        return [
            {
                "creditId": 201,
                "scoreId": 301,
                "scorename": "美育实践学分",
                "scoretypename": "学分",
                "unitcount": "0.50",
                "num": 20,
                "providenum": 0,
                "finish": False,
            },
            {
                "creditId": 202,
                "scoreId": 302,
                "scorename": "劳动教育学分",
                "scoretypename": "学分",
                "unitcount": "0.50",
                "num": 20,
                "providenum": 0,
                "finish": False,
            },
        ]


class PriorityFallbackPlanningDMAPI(FakeDMAPI):
    CREDIT_URL_CREDITEDMEM = "credited"

    def __init__(
        self,
        export_results: list[bytes | None],
        credit_rows: list[dict],
    ):
        super().__init__(export_results)
        self.credit_rows = credit_rows

    def get_mime_manage_activity_dict(self):
        return {
            "1001": {"activityId": 1001, "name": "综合活动A"},
        }

    def get_creditType_list(self, activity_id: str):
        assert str(activity_id) == "1001"
        return self.credit_rows

    def get_credit_list(self, url: str, activity_id: str, score_id: str):
        assert url == self.CREDIT_URL_CREDITEDMEM
        return []


class PreviewPlanningDMAPI(FakeDMAPI):
    CREDIT_URL_CREDITEDMEM = "credited"

    def __init__(
        self,
        export_results: list[bytes | None],
        credited_by_credit: dict[tuple[str, str], set[str]] | None = None,
    ):
        super().__init__(export_results)
        self.credited_by_credit = {
            key: set(values) for key, values in (credited_by_credit or {}).items()
        }
        self.credit_list_calls: list[tuple[str, str]] = []

    def get_mime_manage_activity_dict(self):
        return {
            "1001": {"activityId": 1001, "name": "美育活动A"},
        }

    def get_creditType_list(self, activity_id: str):
        assert str(activity_id) == "1001"
        return [
            {
                "creditId": 201,
                "scoreId": 301,
                "scorename": "美育实践学分",
                "scoretypename": "学分",
                "unitcount": "0.20",
                "num": 20,
                "providenum": 0,
                "finish": False,
            },
            {
                "creditId": 202,
                "scoreId": 302,
                "scorename": "美育实践学分",
                "scoretypename": "学分",
                "unitcount": "0.30",
                "num": 20,
                "providenum": 0,
                "finish": False,
            },
        ]

    def get_credit_list(self, url: str, activity_id: str, score_id: str):
        assert url == self.CREDIT_URL_CREDITEDMEM
        self.credit_list_calls.append((activity_id, score_id))
        return [
            {"signUpId": sign_up_id}
            for sign_up_id in sorted(
                self.credited_by_credit.get((activity_id, score_id), set())
            )
        ]


def build_input_workbook_file(path, rows: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["学号", "姓名", "学分类型", "学分数值"])
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def test_student_is_eligible_only_when_in_admit_list():
    dmapi = FakeDMAPI(
        [
            build_admit_workbook_bytes(
                [
                    ["9001", "美育活动A", "录取", "张三", "20250001", "u-001"],
                    ["9002", "美育活动A", "录取", "王五", "20250003", "u-003"],
                ]
            )
        ]
    )
    bundle = build_bundle()
    admitted = DemandRecord("20250001", "张三", "美育实践学分", 50)
    excluded = DemandRecord("20250002", "李四", "美育实践学分", 50)

    eligibility = build_eligibility_map(dmapi, [bundle], [admitted, excluded])

    assert eligibility.is_eligible("20250001", "1001", "美育实践学分")
    assert not eligibility.is_eligible("20250002", "1001", "美育实践学分")
    match = eligibility.get_match("20250001", "张三", "美育实践学分", "1001")
    assert match is not None
    assert match.sign_up_id == "9001"
    assert match.user_id == "u-001"


def test_parses_user_id_from_header_when_admit_columns_shift():
    dmapi = FakeDMAPI(
        [
            build_admit_workbook_bytes(
                [
                    [
                        "9001",
                        "美育活动A",
                        "录取",
                        "张三",
                        "20250001",
                        "机械工程学院",
                        "机械12501班",
                        "61057685",
                    ]
                ],
                header=[
                    "报名ID",
                    "活动名称",
                    "状态",
                    "姓名",
                    "学号",
                    "学院",
                    "班级",
                    "UID",
                ],
            )
        ]
    )
    bundle = build_bundle()
    demand = DemandRecord("20250001", "张三", "美育实践学分", 50)

    eligibility = build_eligibility_map(dmapi, [bundle], [demand])
    match = eligibility.get_match("20250001", "张三", "美育实践学分", "1001")

    assert match is not None
    assert match.user_id == "61057685"


def test_collects_not_in_admit_list_for_unmatched_demands():
    dmapi = FakeDMAPI(
        [
            build_admit_workbook_bytes(
                [["9001", "美育活动A", "录取", "张三", "20250001", "u-001"]]
            )
        ]
    )
    bundle = build_bundle()
    mismatched_name = DemandRecord("20250001", "李四", "美育实践学分", 50)

    eligibility = build_eligibility_map(dmapi, [bundle], [mismatched_name])

    assert eligibility.not_in_admit_list == (mismatched_name,)


def test_retries_admit_export_until_success():
    success_payload = build_admit_workbook_bytes(
        [["9001", "美育活动A", "录取", "张三", "20250001", "u-001"]]
    )
    dmapi = FakeDMAPI([None, None, success_payload])
    bundle = build_bundle()
    demand = DemandRecord("20250001", "张三", "美育实践学分", 50)

    eligibility = build_eligibility_map(dmapi, [bundle], [demand], retry_attempts=3)

    assert eligibility.is_eligible("20250001", "1001", "美育实践学分")
    assert dmapi.export_calls == [("1001", FakeDMAPI.EXPORT_TYPE_ADMIT)] * 3


def test_does_not_retry_deterministic_admit_sheet_validation_failures():
    malformed_payload = build_admit_workbook_bytes(
        [["9001", "美育活动A", "录取", "张三", None, "u-001"]]
    )
    success_payload = build_admit_workbook_bytes(
        [["9001", "美育活动A", "录取", "张三", "20250001", "u-001"]]
    )
    dmapi = FakeDMAPI([malformed_payload, success_payload, success_payload])
    bundle = build_bundle()
    demand = DemandRecord("20250001", "张三", "美育实践学分", 50)

    try:
        build_eligibility_map(dmapi, [bundle], [demand], retry_attempts=3)
    except ValueError as exc:
        assert str(exc) == "Admit list row is missing required fields"
    else:
        raise AssertionError("Expected deterministic admit-sheet validation failure")

    assert dmapi.export_calls == [("1001", FakeDMAPI.EXPORT_TYPE_ADMIT)]


def test_plan_workflow_writes_plan_json_and_preview_reports(tmp_path):
    excel_path = tmp_path / "input.xlsx"
    build_input_workbook_file(
        excel_path,
        [
            ["20250001", "张三", "美育实践学分", "0.50"],
            ["20250002", "李四", "美育实践学分", "0.50"],
        ],
    )
    dmapi = FakePlanningDMAPI(
        [
            build_admit_workbook_bytes([["9001", "美育活动A", "录取", "张三", "20250001", "u-001"]]),
            build_admit_workbook_bytes([]),
        ]
    )

    result = run_plan_workflow(dmapi=dmapi, excel_path=excel_path, output_dir=tmp_path)

    plan_json_path = result.batch_dir / "plan.json"
    assert plan_json_path.exists()
    assert (result.batch_dir / "plan_summary.xlsx").exists()
    assert (result.batch_dir / "activities.xlsx").exists()
    assert (result.batch_dir / "not_in_admit_list.xlsx").exists()
    assert (result.batch_dir / "activities" / "1001_美育活动A" / "not_in_admit_list.xlsx").exists()
    assert (result.batch_dir / "activities" / "1002_美育活动B" / "not_in_admit_list.xlsx").exists()

    plan_payload = json.loads(plan_json_path.read_text(encoding="utf-8"))
    assert plan_payload["allocations"]
    assert plan_payload["not_in_admit_list"][0]["student_id"] == "20250002"


def test_plan_workflow_uses_excel_name_and_timestamp_for_default_batch_dir(tmp_path):
    excel_path = tmp_path / "王者.xlsx"
    build_input_workbook_file(
        excel_path,
        [
            ["20250001", "张三", "美育实践学分", "0.50"],
        ],
    )
    dmapi = FakePlanningDMAPI(
        [
            build_admit_workbook_bytes(
                [["9001", "美育活动A", "录取", "张三", "20250001", "u-001"]]
            ),
            build_admit_workbook_bytes([]),
        ]
    )

    result = run_plan_workflow(dmapi=dmapi, excel_path=excel_path, output_dir=tmp_path)

    assert result.batch_dir.parent == tmp_path
    assert re.fullmatch(r"王者_\d{4}-\d{2}-\d{2}_\d{6}", result.batch_dir.name)


def test_plan_workflow_excludes_activities_without_sign_card(tmp_path):
    excel_path = tmp_path / "input.xlsx"
    build_input_workbook_file(
        excel_path,
        [
            ["20250001", "张三", "美育实践学分", "0.50"],
            ["20250002", "李四", "美育实践学分", "0.50"],
        ],
    )
    dmapi = SignCardPlanningDMAPI(
        [
            build_admit_workbook_bytes(
                [["9001", "美育活动A", "录取", "张三", "20250001", "u-001"]]
            ),
        ],
        sign_ids={"1001": "sign-1001", "1002": None},
    )

    result = run_plan_workflow(dmapi=dmapi, excel_path=excel_path, output_dir=tmp_path)

    assert dmapi.sign_id_calls == ["1001", "1002"]
    assert dmapi.credit_type_calls == ["1001"]
    assert dmapi.export_calls == [("1001", FakeDMAPI.EXPORT_TYPE_ADMIT)]
    assert {bundle["activity_id"] for bundle in result.plan_payload["bundles"]} == {"1001"}
    assert {
        assignment["bundle"]["activity_id"]
        for allocation in result.plan_payload["allocations"]
        for assignment in allocation["assignments"]
    } == {"1001"}
    assert result.plan_payload["not_in_admit_list"][0]["student_id"] == "20250002"


def test_plan_workflow_keeps_all_credit_types_in_per_activity_exception_workbook(tmp_path):
    excel_path = tmp_path / "input.xlsx"
    build_input_workbook_file(
        excel_path,
        [
            ["20250001", "张三", "美育实践学分", "0.50"],
            ["20250002", "李四", "劳动教育学分", "0.50"],
        ],
    )
    dmapi = FakeMultiBundlePlanningDMAPI(
        [
            build_admit_workbook_bytes([]),
        ]
    )

    result = run_plan_workflow(dmapi=dmapi, excel_path=excel_path, output_dir=tmp_path)

    workbook = load_workbook(
        result.batch_dir / "activities" / "1001_综合活动A" / "not_in_admit_list.xlsx"
    )
    rows = list(workbook.active.iter_rows(values_only=True))

    assert rows == [
        ("学号", "姓名", "学分类型", "应发值(cent)"),
        ("20250001", "张三", "美育实践学分", 50),
        ("20250002", "李四", "劳动教育学分", 50),
    ]


def test_plan_workflow_reports_partial_and_full_preissued_states(tmp_path):
    excel_path = tmp_path / "input.xlsx"
    build_input_workbook_file(
        excel_path,
        [
            ["20250001", "张三", "美育实践学分", "0.50"],
            ["20250002", "李四", "美育实践学分", "0.50"],
        ],
    )
    dmapi = PreviewPlanningDMAPI(
        [
            build_admit_workbook_bytes(
                [
                    ["9001", "美育活动A", "录取", "张三", "20250001", "u-001"],
                    ["9002", "美育活动A", "录取", "李四", "20250002", "u-002"],
                ]
            )
        ],
        credited_by_credit={
            ("1001", "201"): {"9001", "9002"},
            ("1001", "202"): {"9002"},
        },
    )

    result = run_plan_workflow(dmapi=dmapi, excel_path=excel_path, output_dir=tmp_path)

    partial_rows = list(
        load_workbook(result.batch_dir / "already_partially_issued.xlsx")
        .active.iter_rows(values_only=True)
    )
    full_rows = list(
        load_workbook(result.batch_dir / "already_fully_issued.xlsx")
        .active.iter_rows(values_only=True)
    )

    assert partial_rows == [
        ("学号", "姓名", "学分类型", "说明"),
        ("20250001", "张三", "美育实践学分", "Some bundle credits were already issued"),
    ]
    assert full_rows == [
        ("学号", "姓名", "学分类型", "说明"),
        ("20250002", "李四", "美育实践学分", "All bundle credits were already issued"),
    ]


def test_plan_workflow_uses_20_plus_30_for_target_40_instead_of_duplicate_20(tmp_path):
    excel_path = tmp_path / "input.xlsx"
    build_input_workbook_file(
        excel_path,
        [
            ["20250001", "张三", "美育实践学分", "0.40"],
        ],
    )
    dmapi = FakePlanningDMAPI(
        [
            build_admit_workbook_bytes(
                [["9001", "美育活动A", "录取", "张三", "20250001", "u-001"]]
            ),
            build_admit_workbook_bytes([]),
        ]
    )

    result = run_plan_workflow(dmapi=dmapi, excel_path=excel_path, output_dir=tmp_path)
    allocation = result.plan_payload["allocations"][0]

    assert allocation["planned_value_cent"] == 50
    assert allocation["delta_cent"] == 10
    assert sorted(
        assignment["bundle"]["bundle_value_cent"]
        for assignment in allocation["assignments"]
    ) == [20, 30]
    assert sorted(
        assignment["bundle"]["credit_items"][0]["credit_id"]
        for assignment in allocation["assignments"]
    ) == ["201", "202"]


def test_plan_workflow_fills_priority_credit_type_shortage_with_next_type(tmp_path):
    excel_path = tmp_path / "input.xlsx"
    build_input_workbook_file(
        excel_path,
        [
            ["20250001", "张三", "劳动教育学分，美育实践学分", "0.50"],
        ],
    )
    dmapi = PriorityFallbackPlanningDMAPI(
        [
            build_admit_workbook_bytes(
                [["9001", "综合活动A", "录取", "张三", "20250001", "u-001"]]
            ),
        ],
        credit_rows=[
            {
                "creditId": 201,
                "scoreId": 301,
                "scorename": "劳动教育学分",
                "scoretypename": "学分",
                "unitcount": "0.30",
                "num": 20,
                "providenum": 0,
                "finish": False,
            },
            {
                "creditId": 202,
                "scoreId": 302,
                "scorename": "美育实践学分",
                "scoretypename": "学分",
                "unitcount": "0.20",
                "num": 20,
                "providenum": 0,
                "finish": False,
            },
        ],
    )

    result = run_plan_workflow(dmapi=dmapi, excel_path=excel_path, output_dir=tmp_path)

    allocation_rows = [
        (
            allocation["demand"]["credit_type"],
            allocation["demand"]["requested_value_cent"],
            allocation["planned_value_cent"],
            allocation["delta_cent"],
        )
        for allocation in result.plan_payload["allocations"]
    ]
    assert allocation_rows == [
        ("劳动教育学分", 30, 30, 0),
        ("美育实践学分", 20, 20, 0),
    ]
    assert [
        (demand["credit_type"], demand["requested_value_cent"])
        for demand in result.plan_payload["demands"]
    ] == [("劳动教育学分", 30), ("美育实践学分", 20)]


def test_plan_workflow_does_not_report_unavailable_earlier_priority_if_later_fills(
    tmp_path,
):
    excel_path = tmp_path / "input.xlsx"
    build_input_workbook_file(
        excel_path,
        [
            ["20250001", "张三", "劳动教育学分，美育实践学分", "0.50"],
        ],
    )
    dmapi = PriorityFallbackPlanningDMAPI(
        [
            build_admit_workbook_bytes(
                [["9001", "综合活动A", "录取", "张三", "20250001", "u-001"]]
            ),
        ],
        credit_rows=[
            {
                "creditId": 202,
                "scoreId": 302,
                "scorename": "美育实践学分",
                "scoretypename": "学分",
                "unitcount": "0.50",
                "num": 20,
                "providenum": 0,
                "finish": False,
            },
        ],
    )

    result = run_plan_workflow(dmapi=dmapi, excel_path=excel_path, output_dir=tmp_path)

    assert result.plan_payload["not_in_admit_list"] == []
    assert [
        allocation["demand"]["credit_type"]
        for allocation in result.plan_payload["allocations"]
    ] == ["美育实践学分"]
    assert result.plan_payload["allocations"][0]["planned_value_cent"] == 50


def test_main_plan_command_invokes_plan_workflow(monkeypatch, tmp_path):
    excel_path = tmp_path / "input.xlsx"
    excel_path.write_text("placeholder", encoding="utf-8")
    observed = {}

    class StubDMAPI:
        def __init__(self, config_path: str):
            observed["config_path"] = config_path

    def fake_run_plan_workflow(dmapi, excel_path_arg, output_dir):
        observed["dmapi"] = dmapi
        observed["excel_path"] = Path(excel_path_arg)
        observed["output_dir"] = Path(output_dir)
        return SimpleNamespace(batch_dir=Path(output_dir) / "batch")

    monkeypatch.setattr(cli_main, "DMAPI", StubDMAPI)
    monkeypatch.setattr(cli_main, "run_plan_workflow", fake_run_plan_workflow)

    exit_code = cli_main.main(
        ["plan", "--excel", str(excel_path), "--output", str(tmp_path)]
    )

    assert exit_code == 0
    assert observed["config_path"] == "config.ini"
    assert observed["excel_path"] == excel_path
    assert observed["output_dir"] == tmp_path


def test_main_execute_requires_exact_confirmation(monkeypatch, tmp_path, capsys):
    plan_path = tmp_path / "plan.json"
    plan_path.write_text(
        json.dumps(
            {
                "allocations": [
                    {
                        "assignments": [
                            {
                                "sign_up_id": "9001",
                                "bundle": {
                                    "activity_id": "1001",
                                    "activity_name": "美育活动A",
                                    "credit_type": "美育实践学分",
                                    "credit_items": [],
                                },
                            }
                        ]
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    calls = []

    class StubDMAPI:
        def __init__(self, config_path: str):
            calls.append(("config", config_path))

    def fake_execute_plan(dmapi, plan_source, output_dir=None):
        calls.append(("execute", dmapi, Path(plan_source), output_dir))
        return SimpleNamespace(
            batch_dir=plan_path.parent,
            issue_success_count=1,
            retry_success_count=0,
            admit_conflict_count=0,
            capacity_conflict_count=0,
            final_failure_count=0,
            total_failure_count=0,
        )

    monkeypatch.setattr(cli_main, "DMAPI", StubDMAPI)
    monkeypatch.setattr(cli_main, "execute_plan", fake_execute_plan)
    monkeypatch.setattr("builtins.input", lambda _: "EXECUTE")

    exit_code = cli_main.main(["execute", "--plan", str(plan_path)])
    output = capsys.readouterr().out

    assert exit_code == 0
    assert "EXECUTE" in output
    assert "计划摘要：" in output
    assert "发放成功 1(涉及1人)" in output
    assert "总失败 0" in output
    assert calls[0] == ("config", "config.ini")
    assert calls[1] == ("execute", calls[1][1], plan_path, None)


def test_main_execute_summary_includes_plan_not_in_admit_list(monkeypatch, tmp_path, capsys):
    plan_path = tmp_path / "plan.json"
    plan_path.write_text(
        json.dumps(
            {
                "allocations": [
                    {
                        "assignments": [
                            {
                                "sign_up_id": "9001",
                                "bundle": {
                                    "activity_id": "1001",
                                    "activity_name": "美育活动A",
                                    "credit_type": "美育实践学分",
                                    "credit_items": [],
                                },
                            }
                        ]
                    }
                ],
                "not_in_admit_list": [
                    {
                        "student_id": "2025003022",
                        "student_name": "马定然",
                        "credit_type": "美育实践学分",
                        "requested_value_cent": 40,
                    }
                ]
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    class StubDMAPI:
        def __init__(self, config_path: str):
            assert config_path == "config.ini"

    def fake_execute_plan(dmapi, plan_source, output_dir=None):
        return SimpleNamespace(
            batch_dir=plan_path.parent,
            issue_success_count=0,
            retry_success_count=0,
            admit_conflict_count=0,
            capacity_conflict_count=0,
            final_failure_count=0,
            total_failure_count=0,
        )

    monkeypatch.setattr(cli_main, "DMAPI", StubDMAPI)
    monkeypatch.setattr(cli_main, "execute_plan", fake_execute_plan)
    monkeypatch.setattr("builtins.input", lambda _: "EXECUTE")

    exit_code = cli_main.main(["execute", "--plan", str(plan_path)])
    output = capsys.readouterr().out

    assert exit_code == 0
    assert "不在录取名单 1(计划期1+执行期0，涉及1人)" in output
    assert "总失败 1(条目，涉及1人)" in output


def test_main_execute_aborts_without_exact_confirmation(
    monkeypatch, tmp_path, capsys
):
    plan_path = tmp_path / "plan.json"
    plan_path.write_text(
        json.dumps(
            {
                "allocations": [
                    {
                        "assignments": [
                            {
                                "sign_up_id": "9001",
                                "bundle": {
                                    "activity_id": "1001",
                                    "activity_name": "美育活动A",
                                    "credit_type": "美育实践学分",
                                    "credit_items": [],
                                },
                            }
                        ]
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    executed = {"called": False}

    class StubDMAPI:
        def __init__(self, config_path: str):
            raise AssertionError("DMAPI should not be initialized before confirmation")

    def fake_execute_plan(dmapi, plan_source, output_dir=None):
        executed["called"] = True
        return SimpleNamespace(
            batch_dir=plan_path.parent,
            issue_success_count=0,
            retry_success_count=0,
            admit_conflict_count=0,
            capacity_conflict_count=0,
            final_failure_count=0,
            total_failure_count=0,
        )

    monkeypatch.setattr(cli_main, "DMAPI", StubDMAPI)
    monkeypatch.setattr(cli_main, "execute_plan", fake_execute_plan)
    monkeypatch.setattr("builtins.input", lambda _: "NOPE")

    exit_code = cli_main.main(["execute", "--plan", str(plan_path)])
    output = capsys.readouterr().out

    assert exit_code == 1
    assert "已取消" in output
    assert executed["called"] is False


def test_main_execute_auto_exits_when_no_planned_issue_actions(
    monkeypatch, tmp_path, capsys
):
    plan_path = tmp_path / "plan.json"
    plan_path.write_text(
        json.dumps(
            {
                "allocations": [{"assignments": []}],
                "not_in_admit_list": [
                    {
                        "student_id": "2025003022",
                        "student_name": "马定然",
                        "credit_type": "美育实践学分",
                        "requested_value_cent": 40,
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    called = {"execute": False}

    class StubDMAPI:
        def __init__(self, config_path: str):
            raise AssertionError("DMAPI should not be initialized when no execution is needed")

    def fake_execute_plan(dmapi, plan_source, output_dir=None):
        called["execute"] = True
        return SimpleNamespace()

    monkeypatch.setattr(cli_main, "DMAPI", StubDMAPI)
    monkeypatch.setattr(cli_main, "execute_plan", fake_execute_plan)
    monkeypatch.setattr(
        "builtins.input",
        lambda _: (_ for _ in ()).throw(
            AssertionError("Input should not be requested when no execution is needed")
        ),
    )

    exit_code = cli_main.main(["execute", "--plan", str(plan_path)])
    output = capsys.readouterr().out

    assert exit_code == 0
    assert "计划无需执行：预计执行0次发放，已自动退出。" in output
    assert "输入 EXECUTE 确认开始执行：" not in output
    assert called["execute"] is False


def test_main_random_drain_command_invokes_workflow(monkeypatch, tmp_path):
    observed = {}

    class StubDMAPI:
        def __init__(self, config_path: str):
            observed["config_path"] = config_path

    def fake_run_random_drain_workflow(
        dmapi,
        threshold_percent,
        output_dir,
        seed=None,
        jitter_count=0,
    ):
        observed["dmapi"] = dmapi
        observed["threshold_percent"] = threshold_percent
        observed["output_dir"] = Path(output_dir)
        observed["seed"] = seed
        observed["jitter_count"] = jitter_count
        return SimpleNamespace(
            batch_dir=Path(output_dir) / "batch",
            successes=(1, 2),
            failures=(3,),
            skipped=(),
        )

    monkeypatch.setattr(cli_main, "DMAPI", StubDMAPI)
    monkeypatch.setattr(
        cli_main,
        "run_random_drain_workflow",
        fake_run_random_drain_workflow,
    )

    exit_code = cli_main.main(
        [
            "random-drain",
            "--threshold-percent",
            "90",
            "--jitter-count",
            "5",
            "--output",
            str(tmp_path),
            "--seed",
            "123",
        ]
    )

    assert exit_code == 0
    assert observed["config_path"] == "config.ini"
    assert observed["threshold_percent"] == 90
    assert observed["jitter_count"] == 5
    assert observed["output_dir"] == tmp_path
    assert observed["seed"] == 123
