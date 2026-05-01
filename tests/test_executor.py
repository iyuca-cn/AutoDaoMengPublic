from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook

import pytest

from credit_workflow.executor import _credited_sign_up_ids_with_retry, execute_plan


def build_admit_workbook_bytes(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["报名ID", "活动名称", "状态", "姓名", "学号", "userId"])
    for row in rows:
        sheet.append(row)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_plan_payload(
    *,
    activity_id: str = "1001",
    activity_name: str = "美育活动A",
    credit_type: str = "美育实践学分",
    requested_value_cent: int = 50,
    sign_up_id: str = "9001",
    user_id: str = "u-001",
    student_id: str = "20250001",
    student_name: str = "张三",
    credit_items: list[dict] | None = None,
) -> dict:
    if credit_items is None:
        credit_items = [
            {
                "credit_id": "201",
                "score_id": "301",
                "credit_type": credit_type,
                "unitcount_cent": 20,
                "remaining_capacity": 9,
            },
            {
                "credit_id": "202",
                "score_id": "302",
                "credit_type": credit_type,
                "unitcount_cent": 30,
                "remaining_capacity": 9,
            },
        ]

    bundle_value_cent = sum(item["unitcount_cent"] for item in credit_items)
    bundle = {
        "activity_id": activity_id,
        "activity_name": activity_name,
        "credit_type": credit_type,
        "bundle_value_cent": bundle_value_cent,
        "bundle_capacity": 9,
        "credit_items": credit_items,
    }

    demand = {
        "student_id": student_id,
        "student_name": student_name,
        "credit_type": credit_type,
        "requested_value_cent": requested_value_cent,
    }

    return {
        "generated_at": "2026-03-24T12:00:00",
        "demands": [demand],
        "bundles": [bundle],
        "not_in_admit_list": [],
        "allocations": [
            {
                "demand": demand,
                "planned_value_cent": bundle_value_cent,
                "delta_cent": bundle_value_cent - requested_value_cent,
                "assignments": [
                    {
                        "bundle": bundle,
                        "sign_up_id": sign_up_id,
                        "user_id": user_id,
                    }
                ],
            }
        ],
    }


class FakeExecutorDMAPI:
    EXPORT_TYPE_ADMIT = 2
    CREDIT_URL_CREDITEDMEM = "credited"
    SIGN_TYPE_UNSIGNED = 1
    SIGN_TYPE_SIGNED = 2

    def __init__(
        self,
        *,
        admit_rows_by_activity: dict[str, list[list[object]]],
        unsigned_by_activity: dict[str, set[str]] | None = None,
        signed_by_activity: dict[str, set[str]] | None = None,
        sign_user_by_activity: dict[str, dict[str, str]] | None = None,
        credited_by_credit: dict[tuple[str, str], set[str]] | None = None,
        send_outcomes: dict[tuple[str, str], list[bool]] | None = None,
        credit_read_outcomes: dict[tuple[str, str], list[object]] | None = None,
        sign_read_outcomes: dict[str, list[object]] | None = None,
    ):
        self.admit_rows_by_activity = admit_rows_by_activity
        self.unsigned_by_activity = {
            key: set(values) for key, values in (unsigned_by_activity or {}).items()
        }
        self.signed_by_activity = {
            key: set(values) for key, values in (signed_by_activity or {}).items()
        }
        self.credited_by_credit = {
            key: set(values) for key, values in (credited_by_credit or {}).items()
        }
        self.send_outcomes = {
            key: list(values) for key, values in (send_outcomes or {}).items()
        }
        self.credit_read_outcomes = {
            key: list(values) for key, values in (credit_read_outcomes or {}).items()
        }
        self.sign_read_outcomes = {
            key: list(values) for key, values in (sign_read_outcomes or {}).items()
        }
        self.export_calls: list[tuple[str, int]] = []
        self.credit_list_calls: list[tuple[str, str]] = []
        self.sign_list_calls: list[str] = []
        self.resign_calls: list[tuple[str, tuple[str, ...], bool]] = []
        self.send_calls: list[tuple[str, str, str]] = []
        derived_sign_user_map = {
            activity_id: {str(row[0]): str(row[5]) for row in rows if len(row) > 5}
            for activity_id, rows in admit_rows_by_activity.items()
        }
        if sign_user_by_activity:
            for activity_id, mapping in sign_user_by_activity.items():
                derived_sign_user_map[activity_id] = {
                    str(sign_up_id): str(user_id)
                    for sign_up_id, user_id in mapping.items()
                }
        self.sign_user_by_activity = derived_sign_user_map
        self._user_lookup = {
            activity_id: {
                str(user_id): str(sign_up_id)
                for sign_up_id, user_id in mapping.items()
            }
            for activity_id, mapping in self.sign_user_by_activity.items()
        }

    def export_mem_excel(self, activity_id: str, export_type: int) -> bytes:
        self.export_calls.append((activity_id, export_type))
        return build_admit_workbook_bytes(self.admit_rows_by_activity.get(activity_id, []))

    def get_sign_list(self, activity_id: str, sign_type: int) -> list[dict]:
        self.sign_list_calls.append(activity_id)
        outcomes = self.sign_read_outcomes.get(activity_id, [])
        if outcomes:
            outcome = outcomes.pop(0)
            if isinstance(outcome, Exception):
                raise outcome
            if outcome is None:
                return None
            if outcome == "CURRENT":
                if sign_type == self.SIGN_TYPE_SIGNED:
                    sign_up_ids = self.signed_by_activity.get(activity_id, set())
                else:
                    sign_up_ids = self.unsigned_by_activity.get(activity_id, set())
            else:
                sign_up_ids = outcome
        else:
            if sign_type == self.SIGN_TYPE_SIGNED:
                sign_up_ids = self.signed_by_activity.get(activity_id, set())
            else:
                sign_up_ids = self.unsigned_by_activity.get(activity_id, set())
        return [
            {
                "signUpId": sign_up_id,
                "userId": self.sign_user_by_activity.get(activity_id, {}).get(
                    str(sign_up_id), ""
                ),
            }
            for sign_up_id in sorted(sign_up_ids)
        ]

    def resign(
        self, activity_id: str, sign_up_id_list: list[str] | tuple[str, ...], is_all: bool = False
    ) -> bool:
        self.resign_calls.append((activity_id, tuple(sign_up_id_list), is_all))
        for sign_up_id in sign_up_id_list:
            self.unsigned_by_activity.setdefault(activity_id, set()).discard(str(sign_up_id))
        return True

    def get_credit_list(self, url: str, activity_id: str, score_id: str) -> list[dict]:
        assert url == self.CREDIT_URL_CREDITEDMEM
        self.credit_list_calls.append((activity_id, score_id))
        outcomes = self.credit_read_outcomes.get((activity_id, score_id), [])
        if outcomes:
            outcome = outcomes.pop(0)
            if isinstance(outcome, Exception):
                raise outcome
            if outcome is None:
                return None
            if isinstance(outcome, tuple) and outcome[0] == "RAW":
                return outcome[1]
            if outcome == "CURRENT":
                sign_up_ids = self.credited_by_credit.get((activity_id, score_id), set())
            else:
                sign_up_ids = outcome
        else:
            sign_up_ids = self.credited_by_credit.get((activity_id, score_id), set())
        return [
            {"signUpId": sign_up_id}
            for sign_up_id in sorted(sign_up_ids)
        ]

    def send_credit(self, activity_id: str, score_id: str, user_list: str) -> bool:
        self.send_calls.append((activity_id, score_id, user_list))
        outcomes = self.send_outcomes.setdefault((activity_id, score_id), [True])
        outcome = outcomes.pop(0)
        if not outcome:
            return False

        user_lookup = self._user_lookup[activity_id]
        credited = self.credited_by_credit.setdefault((activity_id, score_id), set())
        for user_id in user_list.split(","):
            credited.add(user_lookup[user_id])
        return True


class NoSignCardExecutorDMAPI(FakeExecutorDMAPI):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.sign_id_calls: list[str] = []

    def get_signid(self, activity_id: str) -> None:
        self.sign_id_calls.append(str(activity_id))
        return None


def test_execute_plan_resigns_unsigned_students_issues_only_missing_credit_and_writes_reports(
    tmp_path,
):
    plan_payload = make_plan_payload()
    dmapi = FakeExecutorDMAPI(
        admit_rows_by_activity={
            "1001": [["9001", "美育活动A", "录取", "张三", "20250001", "u-001"]]
        },
        unsigned_by_activity={"1001": {"9001"}},
        credited_by_credit={("1001", "201"): {"9001"}},
    )

    result = execute_plan(dmapi, plan_payload, tmp_path / "batch")

    assert result.resign_success_count == 1
    assert result.already_partially_issued_count == 1
    assert result.issue_success_count == 1
    assert result.final_failure_count == 0
    assert dmapi.send_calls == [("1001", "202", "u-001")]
    assert (result.batch_dir / "execution_results.xlsx").exists()
    assert (result.batch_dir / "retry_results.xlsx").exists()
    assert (
        result.batch_dir
        / "activities"
        / "1001_美育活动A"
        / "already_partially_issued.xlsx"
    ).exists()


def test_execute_plan_marks_fully_preissued_without_mutation(tmp_path):
    plan_payload = make_plan_payload()
    dmapi = FakeExecutorDMAPI(
        admit_rows_by_activity={
            "1001": [["9001", "美育活动A", "录取", "张三", "20250001", "u-001"]]
        },
        credited_by_credit={
            ("1001", "201"): {"9001"},
            ("1001", "202"): {"9001"},
        },
    )

    result = execute_plan(dmapi, plan_payload, tmp_path / "batch")

    assert result.already_fully_issued_count == 1
    assert result.issue_success_count == 0
    assert dmapi.resign_calls == []
    assert dmapi.send_calls == []


def test_execute_plan_overrides_invalid_plan_user_id_from_sign_list(tmp_path):
    plan_payload = make_plan_payload(
        user_id="机械工程学院",
        credit_items=[
            {
                "credit_id": "201",
                "score_id": "301",
                "credit_type": "美育实践学分",
                "unitcount_cent": 50,
                "remaining_capacity": 9,
            }
        ],
    )
    dmapi = FakeExecutorDMAPI(
        admit_rows_by_activity={
            "1001": [["9001", "美育活动A", "录取", "张三", "20250001", "机械工程学院"]]
        },
        signed_by_activity={"1001": {"9001"}},
        sign_user_by_activity={"1001": {"9001": "61057685"}},
    )

    result = execute_plan(dmapi, plan_payload, tmp_path / "batch")

    assert result.issue_success_count == 1
    assert result.final_failure_count == 0
    assert dmapi.send_calls == [("1001", "201", "61057685")]


def test_execute_plan_writes_final_student_summary_statistics(tmp_path):
    plan_payload = make_plan_payload(
        credit_items=[
            {
                "credit_id": "201",
                "score_id": "301",
                "credit_type": "美育实践学分",
                "unitcount_cent": 20,
                "remaining_capacity": 9,
            },
            {
                "credit_id": "202",
                "score_id": "302",
                "credit_type": "美育实践学分",
                "unitcount_cent": 30,
                "remaining_capacity": 9,
            },
        ]
    )
    dmapi = FakeExecutorDMAPI(
        admit_rows_by_activity={
            "1001": [["9001", "美育活动A", "录取", "张三", "20250001", "u-001"]]
        }
    )

    result = execute_plan(dmapi, plan_payload, tmp_path / "batch")

    rows = list(
        load_workbook(result.batch_dir / "final_student_summary.xlsx")
        .active.iter_rows(values_only=True)
    )

    assert rows == [
        ("学号", "姓名", "学分类型", "应发值(cent)", "实际发放值(cent)", "最终偏差(cent)", "状态"),
        ("20250001", "张三", "美育实践学分", 50, 50, 0, "精确完成"),
    ]


def test_execute_plan_tracks_issue_success_people_separately_from_issue_times(tmp_path):
    plan_payload = make_plan_payload()
    dmapi = FakeExecutorDMAPI(
        admit_rows_by_activity={
            "1001": [["9001", "美育活动A", "录取", "张三", "20250001", "u-001"]]
        },
    )

    result = execute_plan(dmapi, plan_payload, tmp_path / "batch")

    assert result.issue_success_count == 2
    assert result.issue_success_student_count == 1


def test_execute_plan_retries_recoverable_issue_failures(tmp_path):
    plan_payload = make_plan_payload(
        credit_items=[
            {
                "credit_id": "201",
                "score_id": "301",
                "credit_type": "美育实践学分",
                "unitcount_cent": 50,
                "remaining_capacity": 9,
            }
        ]
    )
    dmapi = FakeExecutorDMAPI(
        admit_rows_by_activity={
            "1001": [["9001", "美育活动A", "录取", "张三", "20250001", "u-001"]]
        },
        send_outcomes={("1001", "201"): [False, True]},
    )

    result = execute_plan(dmapi, plan_payload, tmp_path / "batch", retry_attempts=2)

    assert result.issue_success_count == 1
    assert result.retry_success_count == 1
    assert result.final_failure_count == 0


def test_execute_plan_retries_sign_list_read_failures_instead_of_treating_them_as_empty(
    tmp_path,
):
    plan_payload = make_plan_payload(
        credit_items=[
            {
                "credit_id": "201",
                "score_id": "301",
                "credit_type": "美育实践学分",
                "unitcount_cent": 50,
                "remaining_capacity": 9,
            }
        ]
    )
    dmapi = FakeExecutorDMAPI(
        admit_rows_by_activity={
            "1001": [["9001", "美育活动A", "录取", "张三", "20250001", "u-001"]]
        },
        unsigned_by_activity={"1001": {"9001"}},
        sign_read_outcomes={"1001": [None, "CURRENT", "CURRENT"]},
    )

    result = execute_plan(dmapi, plan_payload, tmp_path / "batch", retry_attempts=2)

    assert result.resign_success_count == 1
    assert result.final_failure_count == 0
    assert dmapi.resign_calls == [("1001", ("9001",), False)]
    assert dmapi.sign_list_calls == ["1001", "1001", "1001"]


def test_execute_plan_skips_activity_without_sign_card(tmp_path):
    plan_payload = make_plan_payload(
        credit_items=[
            {
                "credit_id": "201",
                "score_id": "301",
                "credit_type": "美育实践学分",
                "unitcount_cent": 50,
                "remaining_capacity": 9,
            }
        ]
    )
    dmapi = NoSignCardExecutorDMAPI(
        admit_rows_by_activity={
            "1001": [["9001", "美育活动A", "录取", "张三", "20250001", "u-001"]]
        }
    )

    result = execute_plan(dmapi, plan_payload, tmp_path / "batch")

    assert result.final_failure_count == 1
    assert result.final_failures[0]["note"] == "Activity has no sign card; skipped"
    assert dmapi.sign_id_calls == ["1001"]
    assert dmapi.export_calls == []
    assert dmapi.credit_list_calls == []
    assert dmapi.sign_list_calls == []
    assert dmapi.resign_calls == []
    assert dmapi.send_calls == []


def test_execute_plan_records_capacity_state_conflicts_separately(tmp_path):
    plan_payload = make_plan_payload(
        credit_items=[
            {
                "credit_id": "201",
                "score_id": "301",
                "credit_type": "美育实践学分",
                "unitcount_cent": 50,
                "remaining_capacity": 9,
            }
        ]
    )
    dmapi = FakeExecutorDMAPI(
        admit_rows_by_activity={
            "1001": [["9001", "美育活动A", "录取", "张三", "20250001", "u-001"]]
        },
        send_outcomes={("1001", "201"): [False, False]},
        credit_read_outcomes={("1001", "201"): [set(), set(), set()]},
    )

    result = execute_plan(dmapi, plan_payload, tmp_path / "batch", retry_attempts=2)

    assert result.capacity_conflict_count == 1
    assert result.final_failure_count == 0
    assert result.total_failure_count == 1


def test_execute_plan_retries_post_issue_verification_read_failures(tmp_path):
    plan_payload = make_plan_payload(
        credit_items=[
            {
                "credit_id": "201",
                "score_id": "301",
                "credit_type": "美育实践学分",
                "unitcount_cent": 50,
                "remaining_capacity": 9,
            }
        ]
    )
    dmapi = FakeExecutorDMAPI(
        admit_rows_by_activity={
            "1001": [["9001", "美育活动A", "录取", "张三", "20250001", "u-001"]]
        },
        send_outcomes={("1001", "201"): [True]},
        credit_read_outcomes={("1001", "201"): [set(), None, "CURRENT"]},
    )

    result = execute_plan(dmapi, plan_payload, tmp_path / "batch", retry_attempts=2)

    assert result.issue_success_count == 1
    assert result.retry_success_count == 1
    assert result.final_failure_count == 0
    assert dmapi.credit_list_calls == [("1001", "201"), ("1001", "201"), ("1001", "201")]


def test_execute_plan_records_final_failures_after_retry_budget_exhausted(tmp_path):
    plan_payload = make_plan_payload(
        credit_items=[
            {
                "credit_id": "201",
                "score_id": "301",
                "credit_type": "美育实践学分",
                "unitcount_cent": 50,
                "remaining_capacity": 9,
            }
        ]
    )
    dmapi = FakeExecutorDMAPI(
        admit_rows_by_activity={
            "1001": [["9001", "美育活动A", "录取", "张三", "20250001", "u-001"]]
        },
        send_outcomes={("1001", "201"): [True]},
        credit_read_outcomes={("1001", "201"): [set(), None, None]},
    )

    result = execute_plan(dmapi, plan_payload, tmp_path / "batch", retry_attempts=2)

    assert result.retry_success_count == 0
    assert result.final_failure_count == 1


def test_execute_plan_records_execution_time_admit_conflicts_separately(tmp_path):
    plan_payload = make_plan_payload()
    dmapi = FakeExecutorDMAPI(admit_rows_by_activity={"1001": []})

    result = execute_plan(dmapi, plan_payload, tmp_path / "batch")

    assert result.admit_conflict_count == 1
    assert result.final_failure_count == 0
    assert dmapi.resign_calls == []
    assert dmapi.send_calls == []


def test_execute_plan_writes_per_activity_admit_and_capacity_conflict_workbooks(tmp_path):
    admit_conflict_plan = make_plan_payload()
    admit_conflict_dmapi = FakeExecutorDMAPI(admit_rows_by_activity={"1001": []})

    admit_result = execute_plan(admit_conflict_dmapi, admit_conflict_plan, tmp_path / "admit")

    assert (
        admit_result.batch_dir
        / "activities"
        / "1001_美育活动A"
        / "admit_conflicts.xlsx"
    ).exists()

    capacity_conflict_plan = make_plan_payload(
        credit_items=[
            {
                "credit_id": "201",
                "score_id": "301",
                "credit_type": "美育实践学分",
                "unitcount_cent": 50,
                "remaining_capacity": 9,
            }
        ]
    )
    capacity_conflict_dmapi = FakeExecutorDMAPI(
        admit_rows_by_activity={
            "1001": [["9001", "美育活动A", "录取", "张三", "20250001", "u-001"]]
        },
        send_outcomes={("1001", "201"): [False, False]},
        credit_read_outcomes={("1001", "201"): [set(), set(), set()]},
    )

    capacity_result = execute_plan(
        capacity_conflict_dmapi,
        capacity_conflict_plan,
        tmp_path / "capacity",
        retry_attempts=2,
    )

    assert (
        capacity_result.batch_dir
        / "activities"
        / "1001_美育活动A"
        / "capacity_conflicts.xlsx"
    ).exists()


def test_credited_sign_up_ids_with_retry_does_not_retry_deterministic_parse_errors():
    dmapi = FakeExecutorDMAPI(
        admit_rows_by_activity={"1001": []},
        credit_read_outcomes={("1001", "201"): [("RAW", [{"bad": "row"}]), "CURRENT"]},
    )

    with pytest.raises(KeyError):
        _credited_sign_up_ids_with_retry(dmapi, "1001", "201", retry_attempts=2)

    assert dmapi.credit_list_calls == [("1001", "201")]
